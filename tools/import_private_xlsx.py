"""One-time private XLSX importer for Bourbon Buddies Club Firestore data.

Run without --apply first. It reads the workbook but does not send data anywhere.
Run with --apply only after reviewing the summary and Firestore rules.
"""

import argparse
import hashlib
import re
import sys
from datetime import date, datetime, time, timezone
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    """Return a trimmed value or None for blank cells."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def utc_datetime(value):
    """Convert an Excel date/datetime to a Firestore-compatible UTC datetime."""
    if isinstance(value, datetime):
        return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value
    if isinstance(value, date):
        return datetime.combine(value, time.min, tzinfo=timezone.utc)
    if isinstance(value, str):
        normalized = re.sub(r"(\d+)(st|nd|rd|th)", r"\1", value.strip(), flags=re.IGNORECASE)
        for pattern in ("%B %d, %Y", "%b. %d, %Y", "%b %d, %Y"):
            try:
                return datetime.strptime(normalized, pattern).replace(tzinfo=timezone.utc)
            except ValueError:
                continue
    return None


def field_map(sheet):
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    return {str(header).strip(): index for index, header in enumerate(headers) if header}


def value(row, fields, name):
    index = fields.get(name)
    return clean(row[index]) if index is not None and index < len(row) else None


def stable_id(prefix, source_row, seed):
    digest = hashlib.sha256(f"{source_row}:{seed}".encode()).hexdigest()[:16]
    return f"{prefix}-{digest}"


def directory_documents(sheet):
    fields = field_map(sheet)
    documents = []
    for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        name = value(row, fields, "Name")
        if not name:
            continue
        email = value(row, fields, "Email")
        documents.append((
            stable_id("member", source_row, email or name),
            {"name": name, "title": value(row, fields, "Title") or "Member", "phone": value(row, fields, "Phone") or "", "email": email or "", "onlyDramsUsername": value(row, fields, "Only Drams Username") or ""},
        ))
    return documents


def accounting_documents(sheet):
    fields = field_map(sheet)
    documents = []
    for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        description = value(row, fields, "Discription") or value(row, fields, "Description")
        amount = value(row, fields, "Amount")
        entry_date = utc_datetime(value(row, fields, "Date"))
        if not description or amount is None or entry_date is None:
            continue
        documents.append((
            f"sheet-row-{source_row}",
            {"date": entry_date, "description": description, "status": value(row, fields, "Status") or "", "member": value(row, fields, "Group member") or "", "amount": float(amount or 0), "notes": value(row, fields, "Notes") or ""},
        ))
    return documents


def newsletter_documents(sheet):
    fields = field_map(sheet)
    link_index = fields.get("Link")
    documents = []
    for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        title = value(row, fields, "Link")
        if not title:
            continue
        link_cell = sheet.cell(row=source_row, column=link_index + 1) if link_index is not None else None
        url = clean(link_cell.hyperlink.target) if link_cell and link_cell.hyperlink else None
        documents.append((f"sheet-row-{source_row}", {"date": utc_datetime(value(row, fields, "Date")), "title": title, "url": url or ""}))
    return documents


def bottle_review_documents(sheet):
    fields = field_map(sheet)
    documents = []
    for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        bottle = value(row, fields, "Bottle (Jack Daniels Single Barell Rye)")
        score = value(row, fields, "Score out of 100")
        if not bottle or score is None:
            continue
        setting = value(row, fields, "Setting (At home, blind in a Glencairn)")
        color = value(row, fields, "Color")
        rating = value(row, fields, "Score Rating System")
        overall_parts = []
        if setting:
            overall_parts.append(f"Setting: {setting}")
        if color:
            overall_parts.append(f"Color: {color}")
        if rating:
            overall_parts.append(f"Original rating note: {rating}")
        documents.append((
            f"original-sheet-row-{source_row}",
            {
                "bottle": bottle,
                "reviewer": value(row, fields, "Reviewed by") or "Original club review",
                "dateReviewed": utc_datetime(value(row, fields, "Date Reviewed")),
                "nose": value(row, fields, "Smelling Notes") or "",
                "palate": value(row, fields, "Tasting Notes") or "",
                "score": float(score),
                "overall": " ".join(overall_parts) or "Imported from the original club workbook.",
                "authorUid": "imported-from-original-xlsx",
                "createdAt": utc_datetime(value(row, fields, "Date Reviewed")) or datetime.now(timezone.utc),
            },
        ))
    return documents


def theme_idea_documents(sheet):
    fields = field_map(sheet)
    documents = []
    for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        theme = value(row, fields, "Theme/topic")
        if not theme:
            continue
        last_used = utc_datetime(value(row, fields, "Last Date Used"))
        documents.append((
            stable_id("original-theme", source_row, theme),
            {
                "theme": theme,
                "lastUsed": last_used,
                "used": last_used is not None,
                "authorUid": "imported-from-original-xlsx",
                "createdAt": last_used or datetime.now(timezone.utc),
            },
        ))
    return documents


def schedule_note_documents(sheet):
    fields = field_map(sheet)
    notes_index = fields.get("Notes")
    documents = []
    for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        event_date = utc_datetime(value(row, fields, "Date"))
        theme = value(row, fields, "Theme")
        notes_cell = sheet.cell(row=source_row, column=notes_index + 1) if notes_index is not None else None
        url = clean(notes_cell.hyperlink.target) if notes_cell and notes_cell.hyperlink else None
        if not event_date or not theme or not url or not url.startswith("https://"):
            continue
        documents.append((
            f"schedule-row-{source_row}",
            {"date": event_date, "theme": theme, "title": value(row, fields, "Notes") or "Tasting Notes", "url": url},
        ))
    return documents


def remove_nones(document):
    return {key: value for key, value in document.items() if value is not None}


def main():
    parser = argparse.ArgumentParser(description="Import private Bourbon Buddies XLSX sheets into Firestore.")
    parser.add_argument("--workbook", required=True, type=Path, help="Path to the original .xlsx workbook")
    parser.add_argument("--service-account", type=Path, help="Firebase Admin SDK JSON key. Required with --apply.")
    parser.add_argument("--apply", action="store_true", help="Write to Firestore. Without this flag, only a summary is shown.")
    args = parser.parse_args()

    if not args.workbook.is_file():
        sys.exit(f"Workbook not found: {args.workbook}")
    if args.apply and not args.service_account:
        sys.exit("--service-account is required with --apply.")
    if args.apply and not args.service_account.is_file():
        sys.exit(f"Service-account key not found: {args.service_account}")

    workbook = load_workbook(args.workbook, data_only=True)
    required = ("Members", "Accounting", "Newsletters", "Bottle Notes", "Tasting Ideas", "Schedule")
    missing = [name for name in required if name not in workbook.sheetnames]
    if missing:
        sys.exit(f"Workbook is missing required sheets: {', '.join(missing)}")

    collections = {
        "privateDirectory": directory_documents(workbook["Members"]),
        "accounting": accounting_documents(workbook["Accounting"]),
        "newsletters": newsletter_documents(workbook["Newsletters"]),
        "bottleReviews": bottle_review_documents(workbook["Bottle Notes"]),
        "themeIdeas": theme_idea_documents(workbook["Tasting Ideas"]),
        "scheduleNotes": schedule_note_documents(workbook["Schedule"]),
    }
    print("Import summary (no private values printed):")
    for name, documents in collections.items():
        print(f"- {name}: {len(documents)} documents")

    if not args.apply:
        print("Dry run complete. Re-run with --apply and a service-account key to write these documents.")
        return

    try:
        import firebase_admin
        from firebase_admin import credentials, firestore
    except ImportError:
        sys.exit("firebase-admin is not installed. Run: py -m pip install firebase-admin openpyxl")

    firebase_admin.initialize_app(credentials.Certificate(str(args.service_account)))
    database = firestore.client()
    batch = database.batch()
    count = 0
    for collection_name, documents in collections.items():
        for document_id, document in documents:
            batch.set(database.collection(collection_name).document(document_id), remove_nones(document), merge=False)
            count += 1
            if count % 400 == 0:
                batch.commit()
                batch = database.batch()
    if count % 400:
        batch.commit()
    print(f"Imported {count} documents. Review them in Firestore, then test with a member account.")


if __name__ == "__main__":
    main()
